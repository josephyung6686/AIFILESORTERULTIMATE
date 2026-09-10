# tests/readers/test_docx_python_docx.py
"""`read_docx` backed by python-docx, against files it actually wrote.

Real `.docx` bytes rather than a stub: the whole point of a reader is what a
library does with a real file, and a fake `read_docx` proves only that the fake
agrees with itself. `65` and `69` both record that every defect this project has
found came from running over real files, and none from the suite.
"""
from __future__ import annotations

import pytest

docx_lib = pytest.importorskip("docx")

from docx.oxml.ns import qn

from readers.docx_python_docx import python_docx_reader


@pytest.fixture
def written(tmp_path):
    """One document carrying every zone this reader claims to tell apart."""
    document = docx_lib.Document()
    document.core_properties.title = "Motion to Compel"
    document.core_properties.author = "Mara Ellison"
    document.add_heading("Background", level=1)
    document.add_paragraph("The deposition was noticed for PHYS 1401.")
    document.add_heading("Argument", level=2)
    document.add_paragraph("Production remains incomplete.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Exhibit A"
    table.cell(1, 1).text = "Withheld"
    path = tmp_path / "motion.docx"
    document.save(path)
    return path


def test_a_real_docx_is_read_into_the_document_p5_expects(written):
    read = python_docx_reader()
    document = read(written)

    assert document is not None, (
        "python-docx is installed and this deployment wired `_no_reader`, so "
        "every .docx on a person's disk recorded `unsupported` -- a missing "
        "library reported where there is none")
    assert document.core_properties["title"] == "Motion to Compel"
    assert document.core_properties["author"] == "Mara Ellison"


def test_headings_are_a_zone_and_body_text_is_not(written):
    document = python_docx_reader()(written)
    zones = {paragraph.text: paragraph.zone for paragraph in document.paragraphs}

    assert zones["Background"] == "heading"
    assert zones["Argument"] == "heading"
    assert zones["The deposition was noticed for PHYS 1401."] == "body"
    assert zones["Production remains incomplete."] == "body"


def test_a_paragraph_carries_the_heading_ancestry_above_it(written):
    document = python_docx_reader()(written)
    by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}

    # A heading paragraph's own last segment IS that heading (P5 `docx.py`).
    assert by_text["Background"].heading_path[-1][1] == "Background"
    # Body under a heading hangs off it.
    assert [label for _, label in
            by_text["The deposition was noticed for PHYS 1401."].heading_path
            ] == ["Background"]
    # A level-2 heading nests under the level-1 above it.
    assert [label for _, label in by_text["Argument"].heading_path] == [
        "Background", "Argument"]


def test_table_cells_are_read_with_their_column_header(written):
    document = python_docx_reader()(written)
    body = {(cell.row, cell.column): cell for cell in document.cells}

    # 1-based throughout: P4 D3 refuses a container-path index of 0.
    assert body[(2, 1)].text == "Exhibit A"
    assert body[(2, 1)].column_header == "Item"
    assert body[(2, 2)].column_header == "Status"
    assert min(cell.row for cell in document.cells) == 1
    assert min(cell.column for cell in document.cells) == 1


def test_paragraph_indexes_are_the_order_they_appear_in(written):
    document = python_docx_reader()(written)
    indexes = [paragraph.index for paragraph in document.paragraphs]

    assert indexes == sorted(indexes), "P4 anchors spans to a paragraph ordinal"
    assert len(set(indexes)) == len(indexes), "two paragraphs cannot share an anchor"
    assert min(indexes) == 1, (
        "P4 D3 refuses a container-path index of 0, and a 0 here made the whole "
        "extraction `failed` -- a real document reported as a damaged one")


# --------------------------------------------------------------------------- #
# headers, footers and comments
# --------------------------------------------------------------------------- #
#
# §2.3's list, in full: *"core properties, all paragraphs in order, heading levels,
# tables and table-cell text, HEADERS AND FOOTERS WHERE FEASIBLE, hyperlinks, document
# relationships, and AVAILABLE REVISION OR COMMENT METADATA."* This reader's own
# docstring named headers and footers as the *"honest next increment"*, and the module
# it feeds has carried the `header_footer` zone and `DocxAnnotation` all along -- both
# reachable and both permanently empty, because nothing produced them.
#
# Measured on the owner's 199-file corpus, 2026-09-05, counting `<w:t>` characters:
# 371,345 live in `word/document.xml` and this reader recovers 100% of them; 3,913
# live in `word/header*.xml`, `word/footer*.xml` and `word/comments.xml` and it
# recovered none. Small in characters and not small in content -- the header of every
# résumé in that corpus is the person's name, address and telephone number, and one
# essay carries 1,189 characters of a supervisor's comments.

@pytest.fixture
def with_running_matter(tmp_path):
    """A document whose header, footer and comment each say something the body
    does not."""
    document = docx_lib.Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Joseph Yung — Curriculum Vitae"
    section.footer.paragraphs[0].text = "Page 1 of 2"
    body = document.add_paragraph("Bioengineering, PHYS 1401.")
    document.add_comment(runs=[body.runs[0]] if body.runs else [],
                         text="Add the vascular graft result here.",
                         author="Mara Ellison", initials="ME")
    path = tmp_path / "cv.docx"
    document.save(path)
    return path


def test_a_running_header_and_footer_are_recovered(with_running_matter):
    """The header of a résumé is who it is about, and it was thrown away.

    `zone` is read from WHERE the paragraph sits -- a header part is a header part --
    which is library knowledge exactly as `Region`'s contract requires, and not a
    guess about the text.
    """
    document = python_docx_reader()(with_running_matter)
    running = [p.text for p in document.paragraphs if p.zone == "header_footer"]
    assert "Joseph Yung — Curriculum Vitae" in running
    assert "Page 1 of 2" in running


def test_the_running_matter_does_not_renumber_the_body(with_running_matter):
    """Paragraph ordinals are ADDRESSES -- P4 D3 -- and every stored citation into a
    body paragraph names one. Header and footer paragraphs are numbered after the
    body, so adding them moves no existing address."""
    document = python_docx_reader()(with_running_matter)
    body = [p for p in document.paragraphs if p.zone == "body"]
    assert [p.index for p in body] == list(range(1, len(body) + 1))
    running = [p for p in document.paragraphs if p.zone == "header_footer"]
    assert min(p.index for p in running) > max(p.index for p in body)


def test_a_header_inherited_from_a_previous_section_is_not_stored_twice(tmp_path):
    """Word LINKS a section's header to the previous one unless it is overridden, so
    a walk over sections reports the same running head once per section. Two
    identical `header_footer` paragraphs would be two observations of one value and,
    worse, two `text_units` under different addresses for one piece of text."""
    document = docx_lib.Document()
    document.sections[0].header.paragraphs[0].text = "One running head"
    document.add_paragraph("Body.")
    document.add_section()
    path = tmp_path / "two-sections.docx"
    document.save(path)

    read = python_docx_reader()(path)
    running = [p.text for p in read.paragraphs if p.zone == "header_footer"]
    assert running.count("One running head") == 1


def test_comments_arrive_as_annotations(with_running_matter):
    """§2.3's "available revision or comment metadata". `DocxAnnotation` existed and
    nothing ever built one, so `extract_docx`'s `annotation` branch was dead code."""
    document = python_docx_reader()(with_running_matter)
    assert [a.text for a in document.annotations] == [
        "Add the vascular graft result here."]


def test_the_annotation_name_is_a_slot_word_and_never_the_authors_name():
    """`name` is a SLOT NAME, and putting a value there would put a person into a
    locator.

    `extract_docx` renders it as `segment("field", label=annotation.name)`, and P4 D7
    defines a field label as *"the format's own slot name, verbatim"* -- which is
    what E2's own fixture says too (`tests/p5/test_p5_docx.py`:
    `DocxAnnotation(name="comment", ...)`). `w:comment/@w:author` is the slot; "Mara
    Ellison" is its value. Filling `name` with the author would have produced the
    locator `annotation:field[Mara Ellison]` -- a person's name inside a citation
    string that the review surface renders.

    THE AUTHOR IS THEREFORE DROPPED, and that is a stated limit rather than a
    silence: `DocxAnnotation` has three fields and none of them is an author, so
    carrying one would mean changing an extractor contract, which is not a reader's
    to change. §2.3's clause is served in part -- the comment TEXT arrives -- and the
    author, along with the tracked insertions and deletions the library does not
    surface, is what remains of it. (This sentence used to say "exactly as links and
    relationships are"; `104` §18.2 gap 17 built those, so the comparison would now
    point at work that is done.)
    """
    from readers.docx_python_docx import _ANNOTATION_SLOT

    assert _ANNOTATION_SLOT == "comment"


# --------------------------------------------------------------------------- #
# `104` §18.2 gap 17, loss (a): hyperlinks and document relationships
# --------------------------------------------------------------------------- #
#
# §2.3 asks for them by name and `extractors/docx.py` has carried `zone="link"` and a
# `relationship`-labelled metadata field since it was written. `DocxDocument.links`
# and `.relationships` defaulted to `()` on every real document, so both arms were
# reachable code no `.docx` could enter -- §18.2's definition of a SILENT LOSS: not a
# feature nobody built, but a reading the extractor already knows how to record and
# the reader never hands it.


@pytest.fixture
def with_a_hyperlink(tmp_path):
    """A document that points at a URL from inside a paragraph.

    Built through the OOXML element because python-docx 1.2 READS `w:hyperlink` and
    offers no API to write one. `part.relate_to(..., is_external=True)` is the
    library's own way to declare the relationship, so the fixture exercises exactly
    the two structures the reader reads.
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = docx_lib.Document()
    document.add_heading("Washington University", level=1)
    paragraph = document.add_paragraph("Apply at ")
    reference = document.part.relate_to("https://admissions.wustl.edu",
                                        RT.HYPERLINK, is_external=True)
    run = paragraph.add_run("the admissions page")
    anchor = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): reference})
    run._r.addprevious(anchor)
    anchor.append(run._r)
    document.add_paragraph("Body text about the essay prompt.")
    path = tmp_path / "application.docx"
    document.save(path)
    return path


def test_a_hyperlink_arrives_with_the_paragraph_it_sits_in(with_a_hyperlink):
    """SABOTAGE: collect the hyperlinks from `document.paragraphs` in a second pass
    instead of inside the `_body_blocks` loop. That property skips everything inside
    a table and loses the layout order, so the ordinal would address a different
    paragraph -- and P4 D3 makes a paragraph ordinal an ADDRESS that stored citations
    already name."""
    document = python_docx_reader()(with_a_hyperlink)
    assert [(link.target, link.paragraph) for link in document.links] == [
        ("https://admissions.wustl.edu", 2)]


def test_a_linked_url_is_not_also_reported_as_a_bare_relationship(with_a_hyperlink):
    """One located value, one reading. Every `w:hyperlink` IS a relationship
    reference, so an unfiltered relationship list reports each linked URL twice --
    once at its paragraph and once at no address at all -- and the unlocated copy is
    strictly the worse of the two.

    SABOTAGE: drop the `linked` set difference in `_read`."""
    document = python_docx_reader()(with_a_hyperlink)
    assert "https://admissions.wustl.edu" not in document.relationships


def test_words_own_internal_parts_are_not_relationships(written):
    """The one that keeps this fix from being noise.

    An empty document python-docx writes declares eight INTERNAL relationships --
    `styles.xml`, `stylesWithEffects.xml`, `settings.xml`, `webSettings.xml`,
    `fontTable.xml`, `theme/theme1.xml`, `numbering.xml`, `customXml/item1.xml`.
    `extractors/docx.py` renders each as a `direct` metadata observation, so keeping
    them would mean eight rows of Word's own plumbing in the evidence table for every
    `.docx` on a person's disk, at the strongest reliability the vocabulary has.
    What §2.3 asks to preserve is what the document POINTS AT.

    SABOTAGE: drop the `is_external` check in `_external_relationships`. `is_external`
    is python-docx's reading of OOXML's `TargetMode`, so nothing here inspects the
    string to decide what kind of thing it names.
    """
    document = python_docx_reader()(written)
    assert document.relationships == ()
    assert document.links == ()


def test_a_document_with_no_running_matter_gains_nothing(written):
    """The guard. A plain document must not acquire an empty header paragraph, an
    empty footer or a phantom annotation: an empty tuple from a reader that looked
    and an empty tuple from a document that has none must stay the same value."""
    document = python_docx_reader()(written)
    assert [p for p in document.paragraphs if p.zone == "header_footer"] == []
    assert document.annotations == ()
